import asyncio
import hashlib
import logging
import os
import re
import shutil
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import numpy as np
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Depends, Header
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from langchain_community.vectorstores import FAISS
from langchain_community.vectorstores.utils import DistanceStrategy
from openpyxl import Workbook
from pdf_ocr import extract_chunks, extract_text, embeddings
from pydantic import BaseModel
from extract import run_batch, spec_summaries, extract_requirements, gpu_lock
from rag import chat_generate
import auth
from auth import AuthExpired, LoginFailed

logging.basicConfig(level=logging.INFO)
log = logging.getLogger("app")
app = FastAPI()
app.add_middleware(   # only needed if the React dev server is opened without the vite proxy
    CORSMiddleware, allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_methods=["*"], allow_headers=["*"],
)

# ---------------- login (ERP accounts) ----------------
# Optional: limit Tender tools (products, bids, Tender AI) to some ERP roles, e.g. TENDER_ROLES=admin,accounts
# Empty = every logged-in ERP user. The ERP assistant is open to every ERP user; the ERP itself enforces
# what each role may do, because tool calls use that user's own ERP token.
TENDER_ROLES = {r.strip() for r in os.getenv("TENDER_ROLES", "").split(",") if r.strip()}


def current_user(authorization: str = Header(default="")) -> dict:
    user = auth.get(authorization.removeprefix("Bearer ").strip())
    if not user:
        raise HTTPException(status_code=401, detail="Please log in again.")
    return user


def tender_user(user: dict = Depends(current_user)) -> dict:
    if TENDER_ROLES and user["role"] not in TENDER_ROLES:
        raise HTTPException(status_code=403, detail="Your role cannot use the tender tools.")
    return user


def public_user(user: dict) -> dict:
    return {
        "username": user["username"],
        "role": user["role"],
        "tender": not TENDER_ROLES or user["role"] in TENDER_ROLES,
    }


class LoginIn(BaseModel):
    username: str
    password: str


@app.post("/auth/login")
async def auth_login(body: LoginIn):
    try:
        res = await auth.login(body.username.strip(), body.password)
    except LoginFailed as e:
        raise HTTPException(status_code=e.status, detail=e.detail)
    return {"token": res["token"], **public_user(auth.get(res["token"]))}


@app.get("/auth/me")
def auth_me(user: dict = Depends(current_user)):
    return public_user(user)


@app.post("/auth/logout")
def auth_logout(authorization: str = Header(default="")):
    auth.drop(authorization.removeprefix("Bearer ").strip())
    return {"ok": True}

DATA = Path(__file__).parent / "data"
PRODUCTS = DATA / "products"
OUTPUTS = DATA / "outputs"
PRODUCTS.mkdir(parents=True, exist_ok=True)
OUTPUTS.mkdir(parents=True, exist_ok=True)

FILE_WORKERS = int(os.getenv("FILE_WORKERS", "3"))   
CTX_CHARS = int(os.getenv("CTX_CHARS", "1200"))        
CHECK_TOKENS = int(os.getenv("CHECK_TOKENS", "64"))

STATUS_RE = re.compile(r"STATUS:\s*(Yes|No|Partial|Not found)", re.I)
SPEC_RE = re.compile(r"SPEC:\s*(.+)", re.I | re.S)
SCORE = {"yes": 1.0, "partial": 0.5, "no": 0.0, "not found": 0.0}

_index_cache: dict = {}
_index_lock = threading.Lock()
_check_cache: dict = {}  


def slug(name: str) -> str:
    s = re.sub(r"[^a-zA-Z0-9_-]+", "_", name.strip()).strip("_").lower()
    if not s:
        raise HTTPException(status_code=400, detail="Invalid product name")
    return s


def load_index(folder: Path) -> FAISS:
    return FAISS.load_local(
        str(folder),
        embeddings,
        allow_dangerous_deserialization=True,
        distance_strategy=DistanceStrategy.MAX_INNER_PRODUCT,
    )


def get_index(name: str) -> FAISS:
    with _index_lock:
        if name not in _index_cache:
            _index_cache[name] = load_index(PRODUCTS / name)
        return _index_cache[name]


def drop_index(name: str):
    with _index_lock:
        _index_cache.pop(name, None)


def product_names() -> list:
    return [p.name for p in PRODUCTS.iterdir() if (p / "index.faiss").exists()]


erp_agent = None   # ERP agent (cards) - started on app startup, optional


@app.on_event("startup")
async def warm_up():
    global erp_agent
    t = time.time()
    for n in product_names():
        get_index(n)
    log.info("preloaded %d product indexes in %.1fs", len(_index_cache), time.time() - t)
    if os.getenv("ERP_AGENT", "1") == "1":
        try:
            from erp_agent import ERPAgent
            agent = ERPAgent()
            await agent.start()
            erp_agent = agent
        except Exception as e:   # tender features must keep working even if ERP/uv is not set up
            log.warning("ERP agent not started: %s", e)


# ---------------- normal AI (tender assistant) ----------------
GENERAL_PROMPT = (
    "You are Kiwi, the assistant of Acer Biomedicals, developed by Acer Biomedicals. Never mention Qwen "
    "or Alibaba. Help with tenders, bids and medical product specifications. If product context is "
    "given, answer only from it and say so when it is not there. "
    "Answer briefly in the same language the user writes in."
)
_chats: dict = {}


class ChatIn(BaseModel):
    message: str
    session_id: str = "default"
    product: str | None = None      # optional: answer from this product's indexed specs


def _sid(user: dict, session_id: str) -> str:
    return f"{user['username']}:{session_id}"      # one user can never read another user's chat


@app.post("/chat")
def chat(body: ChatIn, user: dict = Depends(tender_user)):
    key = _sid(user, body.session_id)
    msgs = _chats.setdefault(key, [{"role": "system", "content": GENERAL_PROMPT}])
    text = body.message
    if body.product:
        name = slug(body.product)
        if name not in product_names():
            raise HTTPException(status_code=404, detail="Product not found")
        docs = get_index(name).similarity_search(body.message, k=3)
        ctx = "\n\n".join(d.page_content for d in docs)[:CTX_CHARS * 2]
        text = f"Product context:\n{ctx}\n\nQuestion: {body.message}"
    msgs.append({"role": "user", "content": text})
    with gpu_lock:
        reply = chat_generate(msgs, None, 512)
    msgs[-1] = {"role": "user", "content": body.message}   # keep history without the big context
    msgs.append({"role": "assistant", "content": reply})
    if len(msgs) > 13:
        _chats[key] = msgs[:1] + msgs[-12:]
    return {"reply": reply}


# ---------------- ERP agent (cards) ----------------
def _agent() -> "object":
    if erp_agent is None:
        raise HTTPException(status_code=503, detail="ERP agent not running (check ERP_* env / tools.py / logs)")
    return erp_agent


@app.post("/erp/chat")
async def erp_chat(body: ChatIn, user: dict = Depends(current_user)):
    agent, sid = _agent(), "erp:" + _sid(user, body.session_id)
    try:
        return await agent.chat(sid, body.message, user["erp_token"])   # {"reply", "approval"}
    except AuthExpired:
        agent.reset(sid)
        auth.drop(user["session"])
        raise HTTPException(status_code=401, detail="ERP session expired. Please log in again.")


MAX_DOC_CHARS = int(os.getenv("MAX_DOC_CHARS", "6000"))
MAX_DOC_FILES = int(os.getenv("MAX_DOC_FILES", "3"))


def _doc_text(f: UploadFile) -> str:
    name = (f.filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        f.file.seek(0)
        wb = load_workbook(f.file, read_only=True, data_only=True)
        rows = [" | ".join("" if c is None else str(c) for c in r) for ws in wb for r in ws.iter_rows(values_only=True)]
        text = "\n".join(r for r in rows if r.strip(" |"))
    else:
        text = extract_text(f)
    return text.strip()[:MAX_DOC_CHARS]


@app.post("/erp/chat/upload")
async def erp_chat_upload(
    message: str = Form(""),
    session_id: str = Form("default"),
    files: list[UploadFile] = File(default=[]),
    user: dict = Depends(current_user),
):
    agent, sid = _agent(), "erp:" + _sid(user, session_id)
    if len(files) > MAX_DOC_FILES:
        raise HTTPException(status_code=400, detail=f"Attach at most {MAX_DOC_FILES} files.")
    docs = []
    for f in files:
        text = await asyncio.to_thread(_doc_text, f)
        if text:
            docs.append((f.filename or "file", text))
    if files and not docs:
        raise HTTPException(status_code=422, detail="Could not read any text from the attached files.")
    msg = message.strip() or "Read the attached document and prepare the matching ERP entry."
    try:
        return await agent.chat(sid, msg, user["erp_token"], docs)
    except AuthExpired:
        agent.reset(sid)
        auth.drop(user["session"])
        raise HTTPException(status_code=401, detail="ERP session expired. Please log in again.")


class ApprovalIn(BaseModel):
    session_id: str = "default"
    approval_id: str
    approved: bool
    id: str | None = None          # edited values from the UI (only these can be changed)
    id2: str | None = None
    data: dict | None = None


@app.post("/erp/approve")
async def erp_approve(body: ApprovalIn, user: dict = Depends(current_user)):
    agent, sid = _agent(), "erp:" + _sid(user, body.session_id)
    edits = {k: v for k, v in (("id", body.id), ("id2", body.id2), ("data", body.data)) if v is not None}
    try:
        return await agent.approve(sid, body.approval_id, body.approved, edits, user["erp_token"])
    except AuthExpired:
        agent.reset(sid)
        auth.drop(user["session"])
        raise HTTPException(status_code=401, detail="ERP session expired. Please log in again.")


@app.post("/erp/reset")
def erp_reset(session_id: str = Form("default"), user: dict = Depends(current_user)):
    if erp_agent:
        erp_agent.reset("erp:" + _sid(user, session_id))
    _chats.pop(_sid(user, session_id), None)
    return {"ok": True}


def search_all(index: FAISS, mat: np.ndarray, k: int):
    """Saari requirements ek hi FAISS call mein. Returns (top1 scores, per-req contexts)."""
    k = max(1, min(k, index.index.ntotal))
    scores, ids = index.index.search(mat, k)
    top1 = np.where(ids[:, 0] >= 0, scores[:, 0], 0.0)
    ctxs = []
    for row in ids:
        docs = []
        for i in row:
            if i < 0:
                continue
            doc = index.docstore.search(index.index_to_docstore_id[int(i)])
            docs.append(doc.page_content)
        ctxs.append("\n\n".join(docs)[:CTX_CHARS])
    return top1, ctxs


def build_check_prompt(req: str, ctx: str) -> str:
    return (
        f"Product context:\n{ctx}\n\nBid requirement: {req}\n\n"
        "Using only the context, decide if the product meets the requirement. "
        "Reply exactly in this format:\n"
        "STATUS: Yes or No or Partial or Not found\n"
        "SPEC: the product's matching specification in one short line"
    )


def parse_check(req: str, out: str) -> dict:
    m = STATUS_RE.search(out)
    s = SPEC_RE.search(out)
    return {
        "requirement": req,
        "status": m.group(1).capitalize() if m else "Not found",
        "spec": s.group(1).strip().splitlines()[0] if s else out.strip()[:200],
    }


def _h(s: str) -> str:
    return hashlib.sha1(s.encode("utf-8", "ignore")).hexdigest()


def write_excel(path: Path, ranking: list, sheets: dict):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Rank", "Product", "Similarity", "Match %", "Yes", "Partial", "No/Not found"])
    for i, r in enumerate(ranking, 1):
        ws.append([i, r["product"], r["similarity"], r["match_percent"], r["yes"], r["partial"], r["no"]])
    for name, rows in sheets.items():
        w = wb.create_sheet(name[:31])
        w.append(["Requirement", "Status", "Product specification"])
        for row in rows:
            w.append([row["requirement"], row["status"], row["spec"]])
    wb.save(path)


@app.post("/products", dependencies=[Depends(tender_user)])
def create_or_update_product(
    name: str = Form(...),
    usermanual: list[UploadFile] = File(default=[]),
    broucher: list[UploadFile] = File(default=[]),
):
    files = usermanual + broucher
    if not files:
        raise HTTPException(status_code=400, detail="Upload at least one file")
    folder = PRODUCTS / slug(name)

    with ThreadPoolExecutor(min(FILE_WORKERS, len(files))) as ex:
        chunks = [c for res in ex.map(extract_chunks, files) for c in res]

    texts = spec_summaries(chunks)         
    if not texts:
        raise HTTPException(status_code=422, detail="No specifications extracted")
    if (folder / "index.faiss").exists():
        index = load_index(folder)
        index.add_texts(texts)
    else:
        folder.mkdir(parents=True, exist_ok=True)
        index = FAISS.from_texts(
            texts, embeddings, distance_strategy=DistanceStrategy.MAX_INNER_PRODUCT
        )
    index.save_local(str(folder))
    with _index_lock:
        _index_cache[folder.name] = index
    return {"product": folder.name, "chunks_read": len(chunks), "spec_blocks_indexed": len(texts)}


@app.get("/products", dependencies=[Depends(tender_user)])
def list_products():
    return product_names()


@app.delete("/products/{name}", dependencies=[Depends(tender_user)])
def delete_product(name: str):
    folder = PRODUCTS / slug(name)
    if not folder.exists():
        raise HTTPException(status_code=404, detail="Product not found")
    shutil.rmtree(folder)
    drop_index(folder.name)
    return {"deleted": folder.name}


@app.post("/bid", dependencies=[Depends(tender_user)])
def process_bid(
    bid: UploadFile = File(...),
    top_n: int = Form(4),
    max_requirements: int = Form(40),
    context_k: int = Form(3),
):
    t0 = time.time()
    reqs = extract_requirements(extract_text(bid), max_requirements)
    if not reqs:
        raise HTTPException(status_code=422, detail="No requirements found in bid")
    t1 = time.time()

    names = product_names()
    if not names:
        raise HTTPException(status_code=404, detail="No products uploaded yet")

    mat = np.array(embeddings.embed_documents(reqs), dtype="float32")

    scored, ctx_by_product = [], {}
    for n in names:
        top1, ctxs = search_all(get_index(n), mat, context_k)
        scored.append((n, float(np.mean(top1))))
        ctx_by_product[n] = ctxs
    scored.sort(key=lambda x: x[1], reverse=True)
    top = scored[:top_n]

    items = []   
    for n, _ in top:
        for req, ctx in zip(reqs, ctx_by_product[n]):
            p = build_check_prompt(req, ctx)
            items.append((n, req, p, _h(p)))
    todo = list({k: p for _, _, p, k in items if k not in _check_cache}.items())
    if todo:
        outs = run_batch([p for _, p in todo], CHECK_TOKENS)
        if len(_check_cache) > 50000:
            _check_cache.clear()
        for (k, _), o in zip(todo, outs):
            _check_cache[k] = o
    t2 = time.time()

    sheets = {n: [] for n, _ in top}
    for n, req, _, k in items:
        sheets[n].append(parse_check(req, _check_cache[k]))

    ranking = []
    for name, sim in top:
        rows = sheets[name]
        statuses = [r["status"].lower() for r in rows]
        points = sum(SCORE[s] for s in statuses)
        ranking.append({
            "product": name,
            "similarity": round(sim, 4),
            "match_percent": round(100 * points / len(rows), 1),
            "yes": statuses.count("yes"),
            "partial": statuses.count("partial"),
            "no": statuses.count("no") + statuses.count("not found"),
        })

    ranking.sort(key=lambda r: (r["match_percent"], r["similarity"]), reverse=True)
    filename = f"compliance_{int(time.time())}.xlsx"
    write_excel(OUTPUTS / filename, ranking, sheets)
    log.info(
        "bid done: %d reqs | extract %.1fs | check(%d prompts, %d new) %.1fs | total %.1fs",
        len(reqs), t1 - t0, len(items), len(todo), t2 - t1, time.time() - t0,
    )
    return {
        "best_match": ranking[0]["product"],
        "ranking": ranking,
        "sheets": sheets,
        "download": f"/download/{filename}",
    }


@app.get("/download/{filename}", dependencies=[Depends(tender_user)])
def download(filename: str):
    path = OUTPUTS / Path(filename).name
    if not path.exists():
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(path, filename=path.name)

# ---------------- React UI (built with `npm run build` in abm-ui) ----------------
UI_DIST = Path(__file__).parent / "abm-ui" / "dist"
if UI_DIST.exists():   # keep this LAST so it never shadows the API routes
    app.mount("/", StaticFiles(directory=UI_DIST, html=True), name="ui")